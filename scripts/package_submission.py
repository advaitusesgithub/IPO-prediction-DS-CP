"""Build documents and a portable delivery ZIP from recorded results only."""
from pathlib import Path
import csv
import hashlib
import importlib.metadata
import json
import os
from xml.sax.saxutils import escape
from zipfile import ZipFile, ZIP_DEFLATED
import pandas as pd
from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocColor
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / "submission"
TITLE = "IPO Listing Prediction"
DATE = "24 October 2026"
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".runtime/matplotlib"))
(ROOT / ".runtime/matplotlib").mkdir(parents=True, exist_ok=True)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inputs():
    receipt = json.loads((ROOT / "outputs/part3/evaluation_receipt.json").read_text())
    for key in ["input_sha256", "artifact_sha256"]:
        for name, expected in receipt[key].items():
            if sha(ROOT / name) != expected:
                raise ValueError(f"Evaluation evidence changed: {name}")
    return json.loads((ROOT / "outputs/part3/completion_summary.json").read_text())


def report_blocks(s):
    reg, cap, clf, always = s["regression"], s["cap_baseline"], s["classification"], s["always_apply_baseline"]
    b = []
    def h(t): b.append(("h", t))
    def p(t): b.append(("p", t))
    def table(headers, rows): b.append(("table", [headers, *rows]))
    def figure(path, caption): b.append(("image", (path, caption)))
    h("Abstract")
    p("This project implements a reproducible Indian mainboard IPO workflow with two outputs: listing opening price and a positive-listing-gain Apply/Skip signal. Prepared information is restricted to S0, 09:00 IST on the first bidding day. The cohort contains 253 IPOs first listed in 2023–2025. Chronological validation compares 19 candidates, freezes model choices and then evaluates 103 reserved IPOs. A custom historical web interface uses saved held-out predictions, with separately labeled development and training examples.")
    p(f"The selected price method has reserved MAE INR {reg['mae_inr']:.2f}; announced cap alone has MAE INR {cap['mae_inr']:.2f}. The selected classifier achieves {clf['accuracy']:.2%} accuracy by predicting Apply for every IPO, matching the majority baseline. The 90% aspiration and useful Skip detection were not achieved. The delivered contribution is an auditable workflow and working demonstration, with measured model limitations.")
    h("1. Problem, scope and outputs")
    p("An investor-facing forecast must use information available when the decision could be made. S0 is before bidding starts, so final subscription totals, final issue price and listing outcomes cannot be predictors. The regression target is listing opening price in INR. The classifier target is 1 when opening price exceeds final issue price, and 0 otherwise; equality is non-positive. The two outputs come from separate methods and can disagree.")
    p("Apply means predicted positive listing gain before transaction costs and without allotment assumptions. It is not a validated investment-profit strategy. The implemented interface demonstrates historical IPOs rather than collecting live IPO data. Day-wise stages, expanded anchor reputation, post-listing sell/hold and economic backtests are deferred.")
    h("2. Data sources and prepared cohort")
    p("The prepared inputs combine IPO profiles and announced issue terms, an independent cross-exchange listing cohort/verification snapshot, financial extractions from dated pre-bid offer documents, corroborated lead-manager identities, and earlier Nifty/VIX observations. Source names, reconciliations and hashes are retained in the bundled provenance reports. No new download is required to run the models or app.")
    table(["Item", "Prepared coverage"], [["Mainboard IPOs first listed 2023–2025", "253"], ["Original allowed S0 predictors", "52"], ["Core model predictors", "12"], ["Optional financial predictors", "7"], ["IPOs with usable annual income/PAT", "135 of 253"], ["IPOs with corroborated managers", "222 of 253"], ["Frozen development / reserved / excluded", "149 / 103 / 1"]])
    p("Financial extraction coverage is incomplete and the history of managers starts with this cohort. Original heavy raw PDF archives, SME tables and day-wise subscription sources are not runtime inputs. Historical source-path references are provenance, not dependencies outside this package. Rebuilding the original PDF extraction would require the earlier archives. Original source licensing remains applicable.")
    h("3. Cleaning and validation")
    p("Earlier preparation normalized issuer identities, matched entities using ISIN/symbol/name evidence, reconciled listing dates and original listing venues, standardized units and separated financial reporting periods. Reviewed conflicts include original versus later exchange admissions, aliases, and final issue prices that differ from the announced cap. Financial header order was corrected only where source evidence supported it; reported amounts and genuine outliers were preserved.")
    p("The runtime loader checks the 253 unique IPO IDs, exact table coverage, schema and label consistency, positive finite target prices, gain and classification formulas, timezone and the S0 cutoff. Final issue price remains label metadata. No target is imputed. Numeric missing values are retained until fold-specific preprocessing. Median imputation, training-derived missing indicators and optional scaling are fitted on training rows only; entirely missing training columns are retained with zero fill and an indicator.")
    p("The historical extraction audit records 25 checks; the imported EDA audit records 20 checks. Part 1 adds seven readiness checks. Separate tests exercise leakage rejection, label maturity, preprocessing, relocation, saved-model replay and the app. These gates check implementation consistency; they do not prove all historical publication timestamps.")
    h("4. Exploratory analysis and feature choices")
    p("Development EDA uses 149 IPOs. Positive listing gains occur in approximately 78.5% of them, making an always-Apply comparison necessary. Annual income/PAT are available for 76 development IPOs; complete-case deletion would remove many observations. The target is right-skewed and genuine extremes were retained. Strong predictor redundancy and period-varying associations motivate a compact core rather than unrestricted feature search.")
    table(["Group", "Core predictors"], [["Announced terms", "Band cap; relative band width; lot application cost; log planned offer size"], ["Market context", "Nifty 21-session return; 21-session volatility; VIX close"], ["Recent IPO context", "Prior 90-day listing count; median gain"], ["Manager history", "Verified manager count; prior IPO count; median prior gain"]])
    p("The optional seven-variable financial extension adds annual income, PAT margin, income growth, loss status, income/PAT definition flags and financial age. It was compared only using development validation. The winning specifications use the core set. Announced bands/lots and dated RHP availability carry documented assumptions because exact historical publication timestamps are absent; no unverified anchor feature is admitted.")
    h("5. Chronological evaluation design")
    p("Membership uses the S0 date and label availability, not just listing year. Four expanding 2024 validation folds require training S0 and label availability strictly before the first validation cutoff. Same-cutoff IPOs stay together. Fitting/preprocessing are frozen within a quarter. One December-to-January boundary IPO is excluded because its outcome was unavailable at the final fit cutoff.")
    table(["2024 fold", "Train", "Validation", "Unmatured rows purged"], [["Q1", "59", "22", "0"], ["Q2", "80", "14", "1"], ["Q3", "94", "26", "1"], ["Q4", "121", "28", "0"]])
    p("Each of 90 validation IPOs is scored once per candidate. Settings were frozen before final fitting on 149 IPOs at 6 January 2025, 09:00 IST. Both final pipelines stay fixed for 103 reserved IPOs. Market and historical-context snapshots may incorporate earlier outcomes observable before each later IPO's S0; this is not test-target fitting. Earlier preparation displayed full-cohort aggregate outcomes, so the reserved set is not claimed as unseen external data.")
    h("6. Model development and selection")
    p("Ten regression candidates cover announced cap, training-median premium, Ridge with two alpha values and two restrained Random Forest settings, with core/extended feature comparisons. Nine classifier candidates cover always Apply, logistic regression with/without class balancing and two forest settings, again with both feature sets. There are 76 audited candidate/fold fits. Random state is 42. The probability threshold is fixed at 0.5; no threshold search was performed.")
    p("The premium baseline multiplies the announced cap by the median of training opening-price/announced-cap ratios. Learned price models use a prespecified INR 0.01 physical floor. Regression ranking uses pooled MAE, then RMSE and simplicity. Classification ranking uses accuracy, then balanced accuracy, Brier score and simplicity. A baseline can win. Validation outcomes used for choosing models are optimistic selection estimates, not independent final performance.")
    table(["Selected method", "Development validation"], [["Price: training-median-premium baseline", "MAE INR 94.45; RMSE INR 158.02; MAPE 17.97%"], ["Classifier: core Random Forest", "Accuracy 77.78%; balanced accuracy 50.00%"], ["Always Apply on the 90 validation IPOs", "Accuracy 77.78%"]])
    p("The core forest has 200 trees, maximum depth 3, minimum leaf size 5, sqrt feature sampling and no class weighting. Its decision accuracy ties always Apply; a lower Brier score breaks the prespecified tie. It makes no Skip predictions in validation. The best learned price challenger is retained among the development artifacts, but the baseline wins the primary price criterion.")
    h("7. Reserved evaluation results")
    table(["Method", "MAE (INR)", "RMSE (INR)", "MAPE"], [["Frozen selected price method", f"{reg['mae_inr']:.2f}", f"{reg['rmse_inr']:.2f}", f"{reg['mape_pct']:.2f}%"], ["Announced cap", f"{cap['mae_inr']:.2f}", f"{cap['rmse_inr']:.2f}", f"{cap['mape_pct']:.2f}%"]])
    p("Announced cap has lower reserved MAE and MAPE, while the frozen selected method has lower RMSE. The model is not switched after viewing this result. MAPE is not classification accuracy; 100 minus MAPE is not a probability of correct prediction. Rupee error also depends on IPO price level.")
    figure("figures/part3/price_predictions.png", "Figure 1. Saved predictions for the 103 reserved IPOs; the dashed line represents perfect prediction.")
    table(["Classifier metric", "Frozen Random Forest", "Always Apply"], [["Accuracy", f"{clf['accuracy']:.2%}", f"{always['accuracy']:.2%}"], ["Balanced accuracy", f"{clf['balanced_accuracy']:.2%}", f"{always['balanced_accuracy']:.2%}"], ["Apply precision", f"{clf['precision_apply']:.2%}", f"{always['precision_apply']:.2%}"], ["Apply recall", f"{clf['recall_apply']:.2%}", f"{always['recall_apply']:.2%}"], ["F1 (Apply)", f"{clf['f1_apply']:.4f}", f"{always['f1_apply']:.4f}"], ["Brier score (lower is better)", f"{clf['brier']:.4f}", f"{always['brier']:.4f}"], ["Correct non-positive listings", "0 of 33", "0 of 33"]])
    p("Seventy of 103 IPOs have positive gains. The classifier predicts Apply for all 103: TP=70, FP=33, TN=0, FN=0. Its accuracy therefore equals the majority baseline. A lower Brier score does not establish useful Skip decisions or calibrated investment confidence. The 90% aspiration is not reached.")
    figure("figures/part3/confusion_matrix.png", "Figure 2. Reserved confusion matrix: all 103 decisions are Apply.")
    h("8. Implementation and demonstration")
    p("The stack is Python, pandas, NumPy, scikit-learn, Matplotlib, joblib, Jupyter, a Starlette/Uvicorn API and a custom HTML/CSS/JavaScript frontend. The tested environment is Python 3.14 on Linux/WSL; exact versions are locked. Reusable package modules handle loading, splits, guards, metrics, training and final evaluation. All runtime inputs, outputs and code are inside the deliverable, with paths independent of the enclosing folder name.")
    p("The user enters an IPO name or stock symbol and presses Predict. A table displays estimated listing price, Apply/Skip signal, raw positive-gain probability and announced band. A separate historical IPO box supports browsing. Saved results cover 103 reserved and 90 development validation IPOs; 59 earlier training examples are explicitly labeled retrospective illustrations. One excluded boundary IPO has no prediction. Actual outcomes are loaded separately on request, and model performance is available in a details dialog. The app does not fit a model or search the web.")
    p("Both final pipelines contain exact development training IDs, ordered features and cutoff metadata. Reloaded pipelines reproduced all held-out predictions. The evaluation receipt checks input/code/artifact hashes and prevents silent reevaluation against altered evidence. Automated app tests verify selection, displayed predictions and the actual-outcome toggle.")
    h("9. Reproduction and deliverables")
    p("From Windows Command Prompt in the project directory, run start_app.cmd. From WSL/Bash, run bash start_app.sh. Open the local URL printed by Uvicorn, normally http://localhost:8501. The Windows launcher derives its location and uses a separate native Windows environment, avoiding WSL registration requirements. Stop with Ctrl+C.")
    p("After extraction or renaming, recreate .venv using Python 3.14 and requirements.lock.txt. Virtual environments contain generated absolute paths and are not distributed. Windows app setup uses requirements-app.txt with its separate .venv-win environment; Linux uses the full locked environment. Application code and data paths need no edits. submission/START_HERE.md includes commands. The four executed notebooks, prepared tables, model artifacts, reports, tests and local app are packaged in the ZIP. Runtime caches, the original large PDF archives and the ZIP itself are excluded.")
    h("10. Limitations, conclusions and later extension")
    p("The sample is small, source coverage is incomplete, manager histories are cohort-limited and some publication availability is assumed. One label-verification tier is weaker than direct exchange-row agreement. Quarterly classification scores are unstable: one validation quarter contains only one non-positive listing. Class probabilities are uncalibrated. All-cohort aggregate outcomes were seen during preparation.")
    p("All four implementation parts are delivered, but the predictive objective is only partly supported. The price workflow produces measurable forecasts yet does not beat announced cap on reserved MAE. The classifier fails to establish useful Skip detection or 90% accuracy. These findings remain in the report and demo. Later work can examine additional decision-time information, verified subscription stages and economic targets, followed by a fresh evaluation; the current reserved labels must not become a tuning set for a renewed final-performance claim.")
    if (ROOT/'outputs/refinement/summary.json').exists():
        h("11. Final model audit and new-data refinement")
        p("A final review compared 48 candidates across eight expanding 2024–2025 folds (384 fits, 194 validation IPOs). The already examined 2025 period is now development data. Relative-premium regression, cap-proportional training weights, regularization, Random Forest, Extra Trees, histogram boosting and validation-only threshold tuning were checked. No final issue price or future subscription input was admitted.")
        p("Balanced-accuracy selection alone sacrificed too many positive listings. Before new-test scoring, practical development-only guardrails required positive recall at least 75%, non-positive specificity at least 10%, and accuracy no worse than the development majority rate minus five percentage points. The selected refined classifier is a class-balanced Random Forest with threshold 0.45. The initial unconstrained choice and the guardrail adjustment are preserved.")
        p("Seventeen previously unused early-2026 IPOs form the new chronological test. All 253 original IPOs have mature labels before its first cutoff. Only three small existing source snapshots were imported. Later IPOs with stale market context, missing profiles or unverified bands were excluded by recorded eligibility rules. This is a small source-covered test, not a full-year benchmark; original publication-time assumptions remain.")
        table(["Same 17 new IPOs", "Price MAE (INR)", "Classifier accuracy"], [["Original model fit (149 rows)","68.97","29.41%"],["Unchanged method refit (253 rows)","50.95","29.41%"],["Refined v2 (253 rows)","50.95","29.41%"],["Cap / always-Apply baselines","30.67","29.41%"]])
        p("The 26.1% price MAE reduction versus the old fit is due to more training history, not a superior tuned regressor. The announced cap still performs better. All 17 refined classification decisions are Apply: five true positives and twelve false positives. Brier score improves from 0.4315 to 0.2730, but probabilities remain uncalibrated and no decision-accuracy improvement or useful Skip detection is established. No further tuning occurred after new-test scoring.")
        p("The original classifier requires VIX and manager fields unavailable in the new source set; its fitted imputer handles them. The refined model restricts itself to shared observed definitions. That comparison limitation is documented. The fourth executed notebook and docs/final_model_audit.md retain the complete audit, including rejected alternatives. The app labels the 17 new examples separately from original held-out, validation and retrospective training results. Original artifacts and 2025 scores remain unchanged.")
    h("Evidence and attribution")
    p("Recorded project evidence: data/provenance/bundled_inputs.csv; reports/prior_preparation/source_manifest.csv; reports/prior_preparation/processing_summary.json; reports/eda/eda_summary.json; config/model_search.json; config/selected_models.json; outputs/part2/ comparisons and training audits; outputs/part3/metrics.csv and evaluation_receipt.json; outputs/part4/test_results.xml. Historical source manifests retain original provider/file attribution and hashes. This report introduces no new external data or literature claims.")
    return b


def build_reports(blocks):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = DocInches(.7)
    section.left_margin = section.right_margin = DocInches(.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = DocPt(10)
    normal.paragraph_format.space_after = DocPt(6)
    doc.core_properties.title = TITLE
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.add_heading(TITLE, 0)
    doc.add_paragraph("S0 opening-price and positive-listing-gain prediction\nFinal implementation report | " + DATE)
    doc.sections[0].footer.paragraphs[0].text = TITLE + " | Recorded results and limitations"
    import matplotlib
    fonts = Path(matplotlib.get_data_path()) / "fonts/ttf"
    pdfmetrics.registerFont(TTFont("DejaVu", str(fonts / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(fonts / "DejaVuSans-Bold.ttf")))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BodyLocal", fontName="DejaVu", fontSize=9, leading=13, spaceAfter=7))
    styles.add(ParagraphStyle(name="HeadLocal", fontName="DejaVu-Bold", fontSize=12, leading=16, spaceBefore=12, spaceAfter=6, keepWithNext=True))
    styles.add(ParagraphStyle(name="CellLocal", fontName="DejaVu", fontSize=8, leading=11))
    flow = [Paragraph(TITLE, styles["Title"]), Paragraph("S0 opening-price and positive-listing-gain prediction", styles["BodyLocal"]), Paragraph("Final implementation report | " + DATE, styles["BodyLocal"])]
    markdown = ["# " + TITLE, "", "Final implementation report | " + DATE, ""]
    for kind, value in blocks:
        if kind == "h":
            doc.add_heading(value, 1); flow.append(Paragraph(escape(value), styles["HeadLocal"])); markdown += ["## " + value, ""]
        elif kind == "p":
            doc.add_paragraph(value); flow.append(Paragraph(escape(value), styles["BodyLocal"])); markdown += [value, ""]
        elif kind == "table":
            table = doc.add_table(rows=1, cols=len(value[0])); table.style = "Light Shading Accent 1"
            for cell,text in zip(table.rows[0].cells, value[0]): cell.text = text
            for row in value[1:]:
                for cell,text in zip(table.add_row().cells, row): cell.text = text
            content = [[Paragraph(escape(str(cell)), styles["CellLocal"]) for cell in row] for row in value]
            widths = [510/len(value[0])] * len(value[0])
            pdf_table = Table(content, colWidths=widths, repeatRows=1, hAlign="LEFT")
            pdf_table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#dce7f0")), ("VALIGN",(0,0),(-1,-1),"TOP"), ("GRID",(0,0),(-1,-1),.4,colors.HexColor("#ccd5de")), ("LEFTPADDING",(0,0),(-1,-1),6), ("RIGHTPADDING",(0,0),(-1,-1),6), ("TOPPADDING",(0,0),(-1,-1),5), ("BOTTOMPADDING",(0,0),(-1,-1),5)]))
            flow += [pdf_table, Spacer(1,8)]
            markdown += ["| " + " | ".join(value[0]) + " |", "| " + " | ".join(["---"]*len(value[0])) + " |"]
            markdown += ["| " + " | ".join(row) + " |" for row in value[1:]] + [""]
        else:
            path, caption = value
            doc.add_picture(str(ROOT / path), width=DocInches(5.7)); doc.add_paragraph(caption, "Caption")
            img = Image(str(ROOT/path)); scale = min(470/img.imageWidth, 290/img.imageHeight)
            img.drawWidth = img.imageWidth*scale; img.drawHeight = img.imageHeight*scale
            flow.append(KeepTogether([img, Paragraph(escape(caption), styles["BodyLocal"])]))
            markdown += [f"![{caption}](../{path})", ""]
    doc.save(DEST / "IPO_Listing_Prediction_Report.docx")
    def footer(canvas, document):
        canvas.setFont("DejaVu", 8); canvas.setFillColor(colors.gray)
        canvas.drawString(42, 23, TITLE); canvas.drawRightString(A4[0]-42, 23, str(document.page))
    SimpleDocTemplate(str(DEST / "IPO_Listing_Prediction_Report.pdf"), pagesize=A4,
        rightMargin=42, leftMargin=42, topMargin=38, bottomMargin=38, title=TITLE, author="").build(flow, onFirstPage=footer, onLaterPages=footer)
    (DEST / "report.md").write_text("\n".join(markdown), encoding="utf-8")


def build_slides(s):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333); presentation.slide_height = Inches(7.5)
    reg, cap, clf = s["regression"], s["cap_baseline"], s["classification"]
    def slide(title, points, image=None, note=""):
        page = presentation.slides.add_slide(presentation.slide_layouts[6])
        page.background.fill.solid(); page.background.fill.fore_color.rgb = RGBColor.from_string("F6F8FB")
        def text_box(x,y,w,h,text,size,color="17365D",bold=False):
            shape = page.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
            tf=shape.text_frame; tf.word_wrap=True
            for i,line in enumerate(text.split("\n")):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
                p.font.size=Pt(size); p.font.name="Aptos"; p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
                p.space_after=Pt(16)
            return shape
        text_box(.5,.35,12.2,.9,title,30,bold=True)
        text_box(.6,1.5,6.1 if image else 11.9,5.3,"\n".join(points),21,color="263747")
        if image: page.shapes.add_picture(str(ROOT/image),Inches(7.1), Inches(1.7), width=Inches(5.7))
        text_box(.6,7.07,11.5,.25,TITLE + " | " + str(len(presentation.slides)),10,color="6B7785")
        page.notes_slide.notes_text_frame.text=note or " ".join(points)
    slide(TITLE,["Opening-price forecasts and positive-listing-gain signals", "Indian mainboard IPOs | S0 information only", "Final demonstration and measured results", DATE],note="Introduce the two outputs. Explain that all four implementation parts are delivered, while the original 90% aspiration was not achieved.")
    slide("Problem and information cutoff",["S0: 09:00 IST on the first bidding day", "Regression: listing opening price in INR", "Classification: opening price > final issue price", "Final price and outcomes are labels, never predictors", "Apply is a gain signal before costs and allotment" ])
    slide("Prepared data and provenance",["253 mainboard IPOs first listed in 2023–2025", "52 original S0 predictors; 12-variable core", "7 optional financial variables compared", "Annual income/PAT: 135 IPOs; managers: 222", "Prepared tables, source hashes and reconciliation audits bundled"])
    slide("Cleaning and exploratory analysis",["Normalize issuer identities and reconcile listing venues", "Separate annual/interim financial periods and units", "Retain missing values and genuine outliers", "Fit imputation, indicators and scaling inside each fold", "Development positives: 78.5%; compare always Apply"])
    slide("Time-aware evaluation",["149 development / 103 reserved / 1 excluded IPO", "Four expanding 2024 folds: 90 validation IPOs", "Training sizes: 59, 80, 94, 121", "Labels must mature strictly before each cutoff", "Freeze choices before the 2025 evaluation"],note="The excluded December-to-January row has an immature label at final fit. Earlier aggregate outcomes were seen in preparation; do not describe the reserved set as unseen external data.")
    slide("Small, recorded model comparison",["10 price candidates; 9 classifier candidates", "Baselines, Ridge/logistic regression and restrained forests", "Core vs optional financial extension", "76 candidate/fold fits; fixed random seed 42", "Threshold fixed at 0.5; no threshold search"])
    slide("Development selected methods",["Price: training-median premium × announced cap", "Validation price MAE: INR 94.45", "Classifier: core Random Forest, 200 trees, depth 3", "Validation accuracy: 77.78%, equal to always Apply", "All validation decisions were Apply"],note="The premium ratio uses only training opening prices and announced caps. The forest wins the accuracy/balanced-accuracy tie through lower Brier score; that does not improve decisions.")
    slide("Reserved price forecasts",[f"Selected MAE: INR {reg['mae_inr']:.2f}",f"Announced-cap MAE: INR {cap['mae_inr']:.2f}",f"Selected RMSE: INR {reg['rmse_inr']:.2f}",f"Selected MAPE: {reg['mape_pct']:.2f}%", "Keep the frozen choice despite the baseline result"],"figures/part3/price_predictions.png", "Announced cap has lower MAE/MAPE, while the selected method has lower RMSE. Never call 100 minus MAPE accuracy. Model choice is not revised using reserved outcomes.")
    slide("Reserved Apply/Skip results",[f"Accuracy: {clf['accuracy']:.2%}, equal to always Apply", "Balanced accuracy: 50%; F1: 0.8092", "70 positive / 33 non-positive IPOs", "No Skip predictions; zero true negatives", "Useful Skip detection and 90% accuracy not achieved"],"figures/part3/confusion_matrix.png")
    slide("Historical local demonstration",["Enter an IPO name, then press Predict", "Read price, signal and probability in a results table", "Or select from the historical collection", "View actual outcomes separately on request", "Training illustrations are explicitly labeled"],note="Run start_app.cmd from Windows or bash start_app.sh from WSL. Show name search, historical browsing, the model-details dialog and the actual-outcome toggle. Probabilities are uncalibrated.")
    slide("Reproducibility and verification",["Four executed notebooks with outputs", "Exact dependency lock and folder-independent paths", "Saved pipelines reload and reproduce predictions", "Evaluation receipt checks inputs and recorded artifacts", "Tests cover time order, leakage, relocation and app behavior"])
    if (ROOT/'outputs/refinement/summary.json').exists():
        slide("Final audit: genuine new data",["48 candidates; eight expanding folds; 384 fits", "253 original IPOs now serve as development", "17 unused early-2026 IPOs for a new test", "Shared observed inputs; no stale market backfill", "Choices frozen before new-test results"],note="The original 2025 period was already exposed, so it is no longer called an unseen test. Guardrails and thresholds were selected using development outcomes only.")
        slide("What tuning actually changed",["Same 17 IPOs: price MAE 68.97 → 50.95 INR", "Gain comes from more history, not tuning", "Announced cap still wins: 30.67 INR MAE", "Accuracy stays 29.41%; all 17 signals Apply", "Keep the weak-signal conclusion; no 90% claim"],note="Five positive and twelve non-positive outcomes. The refined Brier score is lower, but decision accuracy and Skip detection did not improve. Original artifacts are preserved; the app labels new audit cases separately.")
    slide("Conclusion and next extension",["Complete workflow, documentation and working demo", "Price forecasts measured against a sensible baseline", "Classifier has not learned useful Skip decisions", "Small sample and source-availability assumptions remain", "Extend information and evaluate afresh; do not tune on this test"],note="Conclude with measured limitations. Day-wise stages, anchors, post-listing decisions and economic backtests remain a later extension, not completed features.")
    presentation.core_properties.title=TITLE; presentation.core_properties.author=""
    presentation.core_properties.last_modified_by=""
    presentation.save(DEST / "IPO_Listing_Prediction_Presentation.pptx")
    return len(presentation.slides)


def bundled_files():
    files=[]
    excluded={".venv", ".venv-win", ".runtime", "__pycache__", ".pytest_cache", ".ipynb_checkpoints", ".git"}
    for parent,dirs,names in os.walk(ROOT):
        dirs[:]=[name for name in dirs if name not in excluded]
        for name in names:
            path=Path(parent)/name
            if path.suffix in {".pyc", ".zip", ".sha256"} or name=="MANIFEST.csv": continue
            files.append(path)
    return sorted(files)


def build_archive():
    files=bundled_files()
    with (DEST/"MANIFEST.csv").open("w",newline="",encoding="utf-8") as stream:
        writer=csv.writer(stream); writer.writerow(["relative_path","bytes","sha256"])
        writer.writerows((str(p.relative_to(ROOT)),p.stat().st_size,sha(p)) for p in files)
    zip_path=DEST/"IPO_Listing_Prediction.zip"
    with ZipFile(zip_path,"w",ZIP_DEFLATED,compresslevel=6) as archive:
        for path in [*files,DEST/"MANIFEST.csv"]:
            archive.write(path,"ipo_listing_prediction/"+str(path.relative_to(ROOT)))
    with ZipFile(zip_path) as archive:
        if archive.testzip() is not None: raise ValueError("ZIP integrity failure")
        for path in files:
            assert hashlib.sha256(archive.read("ipo_listing_prediction/"+str(path.relative_to(ROOT)))).hexdigest()==sha(path)
    (DEST / "IPO_Listing_Prediction.zip.sha256").write_text(sha(zip_path) + "  IPO_Listing_Prediction.zip\n")
    return {"file_count":len(files)+1,"archive_bytes":zip_path.stat().st_size,"archive_sha256":sha(zip_path)}


def main():
    DEST.mkdir(exist_ok=True)
    summary=inputs()
    blocks=report_blocks(summary)
    build_reports(blocks)
    slide_count=build_slides(summary)
    pins=sorted(f"{d.metadata['Name']}=={d.version}" for d in importlib.metadata.distributions())
    (ROOT/"requirements.lock.txt").write_text("\n".join(pins)+"\n")
    archive=build_archive()
    print(json.dumps({"report_formats":["docx","pdf","markdown"],"slides":slide_count,**archive},indent=2))


if __name__=="__main__":
    main()
